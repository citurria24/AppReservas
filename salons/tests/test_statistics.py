from datetime import datetime, timedelta
from io import BytesIO
from unittest.mock import patch

from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from salons.models import Reservation
from salons.tests import test_agenda


class StatisticsTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        test_agenda.AgendaPermissionTests.setUpTestData.__func__(cls)
        cls.start = timezone.make_aware(datetime(2026, 5, 1))
        Reservation.objects.all().delete()
        statuses = ['confirmed', 'completed', 'completed', 'no_show', 'cancelled_client', 'cancelled_salon']
        for index, status in enumerate(statuses):
            r = cls.create_reservation(cls.salon, cls.branch, cls.service, cls.professional, cls.start + timedelta(days=index, hours=10), 'Cliente')
            Reservation.objects.filter(pk=r.pk).update(status=status, email='CLIENTE@Example.test' if index % 2 else 'cliente@example.test', reward_discount_percent=15 if index == 1 else 0)
        cls.other = cls.create_reservation(cls.salon, cls.other_branch_same_salon, cls.service, cls.colleague, cls.start + timedelta(days=8, hours=10), 'Otra')
        Reservation.objects.filter(pk=cls.other.pk).update(status='completed', email='otra@example.test')
        cls.create_reservation(cls.other_salon, cls.foreign_branch, cls.foreign_service, cls.foreign_professional, cls.start + timedelta(hours=10), 'Ajeno')
        cls.create_reservation(cls.salon, cls.branch, cls.service, cls.professional, timezone.make_aware(datetime(2026, 6, 1)), 'Junio')
        # UTC May 1 is still April in Uruguay, so it must be excluded.
        cls.create_reservation(cls.salon, cls.branch, cls.service, cls.professional, cls.start - timedelta(hours=2), 'Abril')

    create_reservation = classmethod(test_agenda.AgendaPermissionTests.create_reservation.__func__)

    def report(self, user=None, export=False, **filters):
        self.client.force_login(user or self.owner)
        return self.client.get(reverse('statistics-export' if export else 'statistics', args=[self.salon.slug]), {'year': 2026, 'month': 5, **filters})

    def workbook(self, response):
        self.assertEqual(response.status_code, 200)
        return load_workbook(BytesIO(b''.join(response.streaming_content)))

    def test_owner_metrics_month_and_tenant_isolation(self):
        response = self.report()
        self.assertEqual(response.status_code, 200)
        s = response.context['report']['summary']
        self.assertEqual(s['total'], 7)
        self.assertEqual(s['attended'], 3)
        self.assertEqual(s['absent'], 1)
        self.assertEqual(s['cancelled_client'], 1)
        self.assertEqual(s['cancelled_salon'], 1)
        self.assertEqual(s['attendance_rate'], .75)
        self.assertAlmostEqual(s['cancellation_rate'], 2 / 7)
        self.assertEqual(s['unique_clients'], 2)
        self.assertEqual(s['discounts'], 1)
        self.assertNotContains(response, 'Ajeno')
        groups = response.context['report']['groups']
        self.assertEqual(sum(g['total'] for g in groups['professional']), 7)
        self.assertEqual(sum(g['total'] for g in groups['service']), 7)
        self.assertEqual([g['total'] for g in groups['branch']], [6, 1])

    def test_branch_filter_correct_and_inactive_history_included(self):
        self.other_branch_same_salon.active = False
        self.other_branch_same_salon.save()
        response = self.report(branch=self.other_branch_same_salon.pk)
        self.assertEqual(response.context['report']['summary']['total'], 1)
        self.assertEqual(response.context['report']['summary']['unique_clients'], 1)
        self.assertEqual(response.context['report']['groups']['professional'][0]['name'], self.colleague.display_name)

    def test_admin_and_hairdresser_denied_report_and_excel(self):
        for user in (self.admin, self.hairdresser):
            for export in (False, True):
                self.assertEqual(self.report(user, export=export).status_code, 403)

    def test_navigation_owner_only(self):
        for user in (self.owner, self.admin, self.hairdresser):
            self.client.force_login(user)
            response = self.client.get(reverse('agenda'))
            if user == self.owner:
                self.assertContains(response, '>Estadísticas<')
            else:
                self.assertNotContains(response, '>Estadísticas<')

    def test_other_owner_denied_report_and_excel(self):
        for export in (False, True):
            self.assertEqual(self.report(self.other_owner, export=export).status_code, 404)

    def test_foreign_branch_and_invalid_period_rejected(self):
        for filters in ({'branch': self.foreign_branch.pk}, {'month': 13}, {'year': 'invalid'}, {'month': ''}, {'year': 9999}):
            for export in (False, True):
                self.assertEqual(self.report(export=export, **filters).status_code, 400)

    def test_empty_report_zero_denominators(self):
        response = self.report(month=7)
        self.assertEqual(response.context['report']['summary']['total'], 0)
        self.assertEqual(response.context['report']['summary']['attendance_rate'], 0)
        self.assertEqual(response.context['report']['summary']['cancellation_rate'], 0)
        self.assertEqual(len(self.workbook(self.report(export=True, month=7))['Reservas']['A']), 1)

    def test_xlsx_content_type_filename_sheets_columns_and_month_scope(self):
        response = self.report(export=True)
        self.assertEqual(response['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertIn('TuTurnoUy_agenda-uno_2026-05.xlsx', response['Content-Disposition'])
        self.assertEqual(response['Cache-Control'], 'private, no-store')
        wb = self.workbook(response)
        self.assertEqual(wb.sheetnames, ['Resumen', 'Reservas', 'Profesionales', 'Servicios', 'Sucursales'])
        rows = list(wb['Reservas'].values)
        self.assertEqual(rows[0], ('Fecha', 'Hora inicio', 'Hora fin', 'Sucursal', 'Profesional', 'Servicio', 'Cliente', 'Email', 'Teléfono', 'Estado', 'Descuento/recompensa aplicada'))
        self.assertEqual(len(rows), 8)
        self.assertEqual({row[0].month for row in rows[1:]}, {5})
        self.assertNotIn('Ajeno Demo', [row[6] for row in rows[1:]])
        self.assertEqual(wb['Reservas'].freeze_panes, 'A2')
        self.assertIsNotNone(wb['Reservas'].auto_filter.ref)
        self.assertEqual(wb['Reservas']['A2'].number_format, 'dd/mm/yyyy')
        summary = dict(list(wb['Resumen'].values)[1:])
        self.assertEqual(summary['Total de reservas'], 7)
        self.assertEqual(summary['Reservas con descuento/recompensa aplicada'], 1)
        self.assertEqual(sum(row[1] for row in list(wb['Profesionales'].values)[1:]), 7)
        self.assertEqual(sum(row[1] for row in list(wb['Servicios'].values)[1:]), 7)
        self.assertEqual(sum(row[1] for row in list(wb['Sucursales'].values)[1:]), 7)

    def test_xlsx_branch_scope(self):
        response = self.report(export=True, branch=self.other_branch_same_salon.pk)
        self.assertIn('_este.xlsx', response['Content-Disposition'])
        wb = self.workbook(response)
        self.assertEqual(wb['Reservas'].max_row, 2)
        self.assertEqual(wb['Reservas']['D2'].value, 'Este')
        self.assertEqual(wb['Profesionales']['A2'].value, self.colleague.display_name)
        self.assertEqual(wb['Sucursales'].max_row, 2)

    def test_excel_customer_text_not_executed_as_formula(self):
        Reservation.objects.filter(pk=self.other.pk).update(first_name='=HYPERLINK("https://example.test")', contact='=1+1')
        wb = self.workbook(self.report(export=True, branch=self.other_branch_same_salon.pk))
        self.assertEqual(wb['Reservas']['G2'].data_type, 's')
        self.assertEqual(wb['Reservas']['I2'].data_type, 's')
        self.assertEqual(wb['Reservas']['I2'].value, '=1+1')

    def test_default_period_and_december_boundary(self):
        self.client.force_login(self.owner)
        with patch('django.utils.timezone.localdate', return_value=self.start.date()):
            response = self.client.get(reverse('statistics', args=[self.salon.slug]))
        self.assertEqual(response.context['report']['summary']['total'], 7)
        self.create_reservation(self.salon, self.branch, self.service, self.professional, timezone.make_aware(datetime(2026, 12, 31, 23)), 'Diciembre')
        self.assertEqual(self.report(month=12).context['report']['summary']['total'], 1)
