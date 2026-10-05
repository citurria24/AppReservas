from collections import Counter
from datetime import datetime
from io import BytesIO

from django import forms
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from .models import Branch, Reservation


class StatisticsFilterForm(forms.Form):
    month = forms.IntegerField(label='Mes', min_value=1, max_value=12)
    year = forms.IntegerField(label='Año', min_value=1900, max_value=9998)
    branch = forms.ModelChoiceField(label='Sucursal', queryset=Branch.objects.none(), required=False, empty_label='Todas las sucursales')

    def __init__(self, *args, salon, **kwargs):
        super().__init__(*args, **kwargs)
        # Historical reports include inactive branches too.
        self.fields['branch'].queryset = Branch.objects.filter(salon=salon).order_by('name')


def monthly_report(salon, year, month, branch=None):
    start = timezone.make_aware(datetime(year, month, 1))
    end = timezone.make_aware(datetime(year + (month == 12), 1 if month == 12 else month + 1, 1))
    qs = Reservation.objects.filter(salon=salon, starts_at__gte=start, starts_at__lt=end)
    if branch:
        qs = qs.filter(branch=branch)
    reservations = list(qs.select_related('branch', 'professional', 'service').order_by('starts_at', 'pk'))
    counts = Counter(r.status for r in reservations)
    S = Reservation.Status
    attended = counts[S.COMPLETED]
    absent = counts[S.NO_SHOW]
    cancelled = counts[S.CANCELLED_CLIENT] + counts[S.CANCELLED_SALON]
    total = len(reservations)
    summary = {
        'total': total, 'attended': attended, 'absent': absent,
        'cancelled_client': counts[S.CANCELLED_CLIENT], 'cancelled_salon': counts[S.CANCELLED_SALON],
        'attendance_rate': attended / (attended + absent) if attended + absent else 0,
        'cancellation_rate': cancelled / total if total else 0,
        'unique_clients': len({r.email.strip().casefold() for r in reservations}),
        'discounts': sum(r.reward_discount_percent > 0 for r in reservations),
    }
    groups = {}
    for dimension in ('branch', 'professional', 'service'):
        grouped = {}
        for r in reservations:
            obj = getattr(r, dimension)
            row = grouped.setdefault(obj.pk, {'name': obj.display_name if dimension == 'professional' else obj.name,
                                             'total': 0, 'attended': 0, 'absent': 0, 'cancelled': 0})
            row['total'] += 1
            row['attended'] += r.status == S.COMPLETED
            row['absent'] += r.status == S.NO_SHOW
            row['cancelled'] += r.status in (S.CANCELLED_CLIENT, S.CANCELLED_SALON)
        groups[dimension] = sorted(grouped.values(), key=lambda row: row['name'].casefold())
    return {'summary': summary, 'reservations': reservations, 'groups': groups, 'year': year, 'month': month, 'salon': salon, 'branch': branch}


def summary_rows(report):
    s = report['summary']
    return [
        ('Período', f"{report['year']:04d}-{report['month']:02d}"),
        ('Peluquería', report['salon'].name),
        ('Sucursal', report['branch'].name if report['branch'] else 'Todas las sucursales'),
        ('Total de reservas', s['total']), ('Atendidas', s['attended']), ('Ausentes', s['absent']),
        ('Canceladas por cliente', s['cancelled_client']), ('Canceladas por local', s['cancelled_salon']),
        ('Porcentaje de asistencia', s['attendance_rate']), ('Porcentaje de cancelación', s['cancellation_rate']),
        ('Clientes únicos', s['unique_clients']), ('Reservas con descuento/recompensa aplicada', s['discounts']),
        ('Cálculo de asistencia', 'Atendidas / (atendidas + ausentes)'),
        ('Cálculo de cancelación', '(Canceladas por cliente + por local) / total'),
    ]


def report_workbook(report):
    wb = Workbook()
    wb.remove(wb.active)

    def sheet(name, headers, rows):
        ws = wb.create_sheet(name)
        ws.append(headers)
        for row in rows:
            ws.append([ILLEGAL_CHARACTERS_RE.sub('', value) if isinstance(value, str) else value for value in row])
        # Customer-supplied text must remain text, including strings starting with '='.
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='245B4A')
            cell.alignment = Alignment(wrap_text=True)
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for column in ws.columns:
            width = min(55, max(14, max(len(str(cell.value or '')) for cell in column) + 2))
            ws.column_dimensions[get_column_letter(column[0].column)].width = width
        return ws

    summary = sheet('Resumen', ['Indicador', 'Valor'], summary_rows(report))
    for row in summary.iter_rows(min_row=2):
        if row[0].value in ('Porcentaje de asistencia', 'Porcentaje de cancelación'):
            row[1].number_format = '0.00%'
    rows = []
    for r in report['reservations']:
        starts = timezone.localtime(r.starts_at)
        ends = timezone.localtime(r.ends_at)
        rows.append([starts.date(), starts.time().replace(tzinfo=None), ends.time().replace(tzinfo=None),
                     r.branch.name, r.professional.display_name, r.service.name,
                     f'{r.first_name} {r.last_name}', r.email, r.contact, r.get_status_display(), r.reward_discount_percent / 100])
    reservations = sheet('Reservas', ['Fecha', 'Hora inicio', 'Hora fin', 'Sucursal', 'Profesional', 'Servicio', 'Cliente', 'Email', 'Teléfono', 'Estado', 'Descuento/recompensa aplicada'], rows)
    for row in reservations.iter_rows(min_row=2):
        row[0].number_format = 'dd/mm/yyyy'
        row[1].number_format = row[2].number_format = 'hh:mm'
        row[10].number_format = '0%'
    for name, dimension in [('Profesionales', 'professional'), ('Servicios', 'service'), ('Sucursales', 'branch')]:
        rows = report['groups'][dimension]
        if dimension == 'service':
            sheet(name, ['Servicio', 'Total de reservas', 'Atendidas'], [[r['name'], r['total'], r['attended']] for r in rows])
        else:
            sheet(name, ['Profesional' if dimension == 'professional' else 'Sucursal', 'Total de reservas', 'Atendidas', 'Ausencias', 'Cancelaciones'],
                  [[r['name'], r['total'], r['attended'], r['absent'], r['cancelled']] for r in rows])
    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream
